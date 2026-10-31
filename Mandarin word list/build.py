"""
Build website kamus + deck belajar + latihan dari satu file Excel.

Input : Mandarin word list/Kosakata_Mandarin.xlsx
          - sheet "Survival 200" : kosakata (sumber Deck Belajar dan Latihan)
          - sheet "Kamus"        : kamus (Indonesia, pinyin, hanzi, contoh kalimat)
          - sheet "Progress"     : rekap hafalan (tidak dipakai build)
Output: docs/index.html (kamus + latihan), docs/data.json (data kamus),
        docs/deck.html + docs/deck-slides.html (Deck Belajar), docs/deck.json (soal latihan)

Kolom dicari berdasarkan nama header, bukan posisi.
"""
import collections
import html as H
import json
import re
import sys
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
XLSX = BASE / "Kosakata_Mandarin.xlsx"
TEMPLATE = BASE / "template.html"
DECK_SHELL = BASE / "deck_shell.html"
DECK_TEMPLATE = BASE / "deck_template.html"
OUT_DIR = BASE.parent / "docs"

# ---------- Kamus (sheet "Kamus") ----------
COL = {
    "no": "No.",
    "id": "Kalimat dalam bahasa Indonesia",
    "py": "arti kalimat atau kata dalam pinyin",
    "hz": "arti kalimat atau kata dalam hanzi",
    "words": "arti perkata dalam Bahasa Indonesia",
    "note": "Keterangan",
    "label": "label",
}
EX_TEMPLATE = ("Contoh {k} - Indonesia", "Contoh {k} - Pinyin", "Contoh {k} - Hanzi")

# ---------- Deck (sheet "Survival 200") ----------
VCOL = {"sit": "Situasi", "id": "Indonesia", "py": "Pinyin", "hz": "Hanzi", "words": "Arti per kata", "note": "Keterangan"}
SIT = ['Salam & sopan santun', 'Perkenalan', 'Tidak mengerti & minta bantuan', 'Ya, tidak, boleh', 'Angka & waktu',
       'Transportasi & arah', 'Hotel', 'Makan & minum', 'Belanja & pembayaran',
       'Pelatihan & kerja', 'Sosial & penutup', 'Kesehatan', 'Darurat']
TAHAP = [('Tahap 1 · Dasar', 'b', SIT[:5]), ('Tahap 2 · Perjalanan', 'r', SIT[5:7]),
         ('Tahap 3 · Makanan dan belanja', 'y', SIT[7:9]),
         ('Tahap 4 · Kerja, sosial, kesehatan, darurat', 'g', SIT[9:])]
DAYS = 30
SKIP_NOTE = re.compile("tertulis|dikoreksi|kolom|dirapikan|diganti|diubah|diseragamkan|dilengkapi|^REVIEW|^D\\d+", re.I)
E = H.escape


def clean(v):
    return "" if v is None else str(v).strip()


def headers_of(ws):
    return {clean(c.value): i for i, c in enumerate(ws[1])}


# ================= KAMUS =================
def build_kamus(ws):
    headers = headers_of(ws)
    missing = [v for v in COL.values() if v not in headers]
    if missing:
        sys.exit("Header sheet Kamus tidak ditemukan:\n  - " + "\n  - ".join(missing))
    ex_cols = []
    for k in (1, 2, 3):
        ex_cols.append([headers.get(t.format(k=k)) for t in EX_TEMPLATE])
    entries = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        get = lambda key: clean(row[headers[COL[key]]])
        ident = get("id")
        if not ident:
            continue
        examples = []
        for i_id, i_py, i_hz in ex_cols:
            if i_id is None:
                continue
            e = [clean(row[i_id]), clean(row[i_py]), clean(row[i_hz])]
            if any(e):
                examples.append(e)
        entries.append({"label": get("label"), "id": ident, "py": get("py"), "hz": get("hz"),
                        "words": get("words"), "note": get("note"), "warn": False, "ex": examples})
    if not entries:
        sys.exit("Sheet Kamus kosong.")
    return entries


# ================= KOSAKATA / DECK =================
def load_groups(ws):
    headers = headers_of(ws)
    missing = [v for v in VCOL.values() if v not in headers]
    if missing:
        sys.exit("Header sheet Survival 200 tidak ditemukan:\n  - " + "\n  - ".join(missing))
    groups = collections.OrderedDict()
    for row in ws.iter_rows(min_row=2, values_only=True):
        g = lambda k: clean(row[headers[VCOL[k]]])
        if not g("hz") or not g("id"):
            continue
        key = (g("sit"), g("id"))
        grp = groups.setdefault(key, {"sit": g("sit"), "id": g("id"), "rows": [], "kets": [], "words": g("words")})
        grp["rows"].append((g("hz"), g("py")))
        if g("note") and not SKIP_NOTE.search(g("note")):
            grp["kets"].append(g("note"))
    return list(groups.values())


def num_value(h):
    D = dict(zip('零一二三四五六七八九', range(10)))
    cur = tot = 0
    for c in h:
        if c in D:
            cur = D[c]
        elif c == '十':
            cur = cur or 1; tot += cur * 10; cur = 0
        elif c == '百':
            tot += cur * 100; cur = 0
        elif c == '千':
            tot += cur * 1000; cur = 0
        elif c == '万':
            tot = (tot + cur) * 10000 if tot + cur else 10000; cur = 0
        elif c == '两':
            cur = 2
    return tot + cur


ANGKA_TAGS = ['Jam, tanggal, dan kata tanya', 'Angka 0–20', 'Puluhan dan ratusan',
              'Ribuan dan puluhan ribu', 'Nama bulan', 'Hari, minggu, dan tahun']


def tag_angka(g):
    h = g["rows"][0][0]
    if '？' in h or '点' in h:
        return ANGKA_TAGS[0], 0
    if h.endswith('月') and len(h) <= 3:
        return ANGKA_TAGS[4], num_value(h[:-1])
    if '星期' in h or '年' in h or h in ('今天', '明天', '后天') or h.startswith('下个'):
        return ANGKA_TAGS[5], 0
    if '、' in h:
        return ANGKA_TAGS[1], 1
    if re.fullmatch('[零一二三四五六七八九十百千万两]+', h):
        v = num_value(h)
        return (ANGKA_TAGS[1], v) if v <= 20 else (ANGKA_TAGS[2], v) if v < 1000 else (ANGKA_TAGS[3], v)
    return ANGKA_TAGS[0], 0


def tag_kerja(g):
    i = g["id"]
    proses = ('/' in i and re.search('[A-Za-z]', i)) or i.startswith(('Stasiun', 'Packaging', 'NG', 'OK'))
    return 'Proses dan station produksi' if proses else 'Kerja dan pelatihan'


def plan_days(groups):
    known = set(SIT)
    extra = sorted({g["sit"] for g in groups} - known)
    sits = SIT + extra
    tahap = list(TAHAP)
    if extra:
        print("PERINGATAN: situasi baru dimasukkan ke Tahap 4:", extra)
        tahap[-1] = (tahap[-1][0], tahap[-1][1], tahap[-1][2] + extra)
    by = collections.OrderedDict((s, [g for g in groups if g["sit"] == s]) for s in sits)
    by = collections.OrderedDict((s, l) for s, l in by.items() if l)
    bt = collections.OrderedDict((t, []) for t in ANGKA_TAGS)
    for g in by.get('Angka & waktu', []):
        t, v = tag_angka(g); bt[t].append((v, g))
    angdays = [(t, [g for _, g in sorted(l, key=lambda x: x[0])]) for t, l in bt.items() if l]
    others = [s for s in by if s != 'Angka & waktu']
    rem = DAYS - len(angdays)
    tot = sum(len(by[s]) for s in others)
    alloc = {s: max(1, int(rem * len(by[s]) / tot)) for s in others}
    while sum(alloc.values()) < rem:
        s = max(others, key=lambda s: len(by[s]) / alloc[s]); alloc[s] += 1
    while sum(alloc.values()) > rem:
        s = max(others, key=lambda s: alloc[s] - len(by[s]) / tot * rem if alloc[s] > 1 else -9); alloc[s] -= 1
    days = []
    for s in by:
        if s == 'Angka & waktu':
            days += [(s, t, l) for t, l in angdays]
            continue
        L = by[s]
        if s == 'Pelatihan & kerja':
            L = sorted(L, key=lambda g: tag_kerja(g) != 'Kerja dan pelatihan')
        n = alloc[s]; sz = -(-len(L) // n)
        for k in range(n):
            ch = L[k * sz:(k + 1) * sz]
            tg = {tag_kerja(g) for g in ch} if s == 'Pelatihan & kerja' else set()
            days.append((s, tg.pop() if len(tg) == 1 else s, ch))
    return [d for d in days if d[2]], tahap


def row_html(g):
    main = g["rows"][0]
    h = f'<tr><td>{E(g["id"])}</td><td><span class="hz">{E(main[0])}</span><span class="py">{E(main[1])}</span>'
    for a, p in g["rows"][1:]:
        h += f'<span class="alt"><i>alt</i> {E(a)} · {E(p)}</span>'
    for k in g["kets"][:1]:
        h += f'<span class="note">{E(k)}</span>'
    return h + '</td></tr>'


def card_html(title, col, items):
    h3 = f'<h3>{E(title)}</h3>' if title else '<h3 style="visibility:hidden" aria-hidden="true">&nbsp;</h3>'
    return (f'<div class="card ac {col}">{h3}<table class="t">'
            + ''.join(row_html(g) for g in items) + '</table></div>')


def build_deck(groups):
    days, tahap = plan_days(groups)
    col_of = {s: c for _, c, l in tahap for s in l}
    tahap_of = {s: t for t, _, l in tahap for s in l}
    nrows = len(groups)
    sl = ['<section class="ctr" data-nofoot><h1>Rencana Belajar Harian Mandarin</h1><p class="sub">30 hari · kosakata gabungan edisi survival &amp; kerja</p><p><span class="pill b">Dasar</span> <span class="pill r">Perjalanan</span> <span class="pill y">Makanan &amp; belanja</span> <span class="pill g">Kerja &amp; darurat</span></p></section>']
    sl.append(f'<section><div class="hd"><h2>Cara memakai deck ini</h2><p class="sub">Satu hari, satu topik kecil</p></div><div class="g3"><div class="card kpi b"><b class="cu">{len(days)}</b><span>hari belajar</span></div><div class="card kpi g"><b class="cu">4</b><span>tahap bertema</span></div><div class="card kpi r"><b class="cu">{nrows}</b><span>frasa &amp; kata</span></div></div><ul class="l" style="margin-top:22px"><li data-s="1"><b>Dengar &amp; ulangi:</b> ucapkan tiap frasa dengan lantang, baca pinyin beserta tanda nadanya</li><li data-s="2"><b>Simulasi:</b> bayangkan situasinya (taksi, restoran, tempat pelatihan)</li><li data-s="3"><b>Tanda <span class="alt" style="display:inline-block"><i>alt</i></span>:</b> ada cara lain untuk mengatakan hal yang sama, dengan hanzi berbeda. Cukup hafal salah satu dulu</li></ul></section>')
    rng = {}
    for n, (s, t, l) in enumerate(days, 1):
        a = tahap_of[s]; rng.setdefault(a, [n, n]); rng[a][1] = n
    chain = ''.join(f'<div><b>{t.split(" · ")[0]}</b><br>{t.split(" · ")[1]}<br><small>Hari {rng[t][0]}–{rng[t][1]}</small></div>'
                    for t, _, _ in tahap if t in rng)
    sl.append(f'<section><div class="hd"><h2>Peta 4 tahap</h2><p class="sub">Dari dasar hingga situasi kerja dan darurat</p></div><div class="chain">{chain}</div></section>')
    for n, (s, t, l) in enumerate(days, 1):
        col = col_of[s]; m = -(-len(l) // 12); sz = -(-len(l) // m)
        for p in range(m):
            it = l[p * sz:(p + 1) * sz]
            part = f' · bagian {p + 1}/{m}' if m > 1 else ''
            h = (len(it) + 1) // 2
            title = f'Hari {n} · {t}'
            body = (f'<div class="g2">{card_html(title, col, it[:h])}{card_html(None, col, it[h:])}</div>'
                    if len(it) > 6 else f'<div class="g1">{card_html(title, col, it)}</div>')
            dense = ' class="dn"' if len(it) >= 8 else ''
            sl.append(f'<section{dense}><div class="hd"><h2>{E(tahap_of[s])}</h2><p class="sub">Hari {n}{part} · {len(it)} frasa</p></div>{body}</section>')
    sl.append('<section class="ctr" data-nofoot><h1>Selamat belajar!</h1><p class="sub">Ucapkan lantang, ulangi, dan bayangkan situasinya. Ni jiayou!</p></section>')
    shell = DECK_SHELL.read_text(encoding="utf-8")
    if shell.count("<!--SLIDES-->") != 1:
        sys.exit("deck_shell.html harus memuat tepat satu penanda <!--SLIDES-->")
    return shell.replace("<!--SLIDES-->", "\n".join(sl)), len(days), build_mobile(days, tahap, nrows)


def build_mobile(days, tahap, nrows):
    col_of = {s: c for _, c, l in tahap for s in l}
    tahap_of = {s: t for t, _, l in tahap for s in l}
    chips = ''.join(f'<a class="{col_of[s]}" href="#h{n}">{n}</a>' for n, (s, t, l) in enumerate(days, 1))
    out = [f'<div class="dm-top"><span>{len(days)} hari · {nrows} kosakata</span><a href="deck-slides.html">Buka versi slide ↗</a></div>',
           f'<nav class="dm-chips" aria-label="Pilih hari">{chips}</nav>']
    last = None
    for n, (s, t, l) in enumerate(days, 1):
        if tahap_of[s] != last:
            last = tahap_of[s]
            out.append(f'<h2 class="dm-t {col_of[s]}">{E(last)}</h2>')
        items = ''
        for g in l:
            main = g["rows"][0]
            items += f'<div class="dm-i"><div class="dm-id">{E(g["id"])}</div><div class="dm-hz">{E(main[0])}</div><div class="dm-py">{E(main[1])}</div>'
            for a, p in g["rows"][1:]:
                items += f'<span class="dm-alt"><i>alt</i> {E(a)} · {E(p)}</span>'
            for k in g["kets"][:1]:
                items += f'<span class="dm-note">{E(k)}</span>'
            items += '</div>'
        out.append(f'<section class="dm-day {col_of[s]}" id="h{n}"><h3>Hari {n} · {E(t)}<small>{len(l)} frasa</small></h3>{items}</section>')
    return "\n".join(out)


def build_pool(groups):
    return [{"label": "", "id": g["id"], "py": g["rows"][0][1], "hz": g["rows"][0][0],
             "words": g["words"], "note": g["kets"][0] if g["kets"] else "", "warn": False, "ex": []}
            for g in groups]


def inject(html, marker, data):
    if html.count(marker) != 1:
        sys.exit(f"template.html harus memuat tepat satu baris '{marker}'")
    return html.replace(marker, marker[:-3] + json.dumps(data, ensure_ascii=False) + ";")


def main():
    if not XLSX.exists():
        sys.exit(f"File tidak ditemukan: {XLSX}")
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    for name in ("Survival 200", "Kamus"):
        if name not in wb.sheetnames:
            sys.exit(f'Sheet "{name}" tidak ditemukan di {XLSX.name}')
    entries = build_kamus(wb["Kamus"])
    groups = load_groups(wb["Survival 200"])
    pool = build_pool(groups)
    deck_html, ndays, mobile = build_deck(groups)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUT_DIR / "data.json").write_text(json.dumps(entries, ensure_ascii=False, indent=1), encoding="utf-8")
    (OUT_DIR / "deck.json").write_text(json.dumps(pool, ensure_ascii=False, indent=1), encoding="utf-8")
    html = TEMPLATE.read_text(encoding="utf-8")
    html = inject(html, "var DATA=[];", entries)
    html = inject(html, "var DECK=[];", pool)
    (OUT_DIR / "index.html").write_text(html, encoding="utf-8")
    (OUT_DIR / "deck-slides.html").write_text(deck_html, encoding="utf-8")
    dt = DECK_TEMPLATE.read_text(encoding="utf-8")
    if dt.count("<!--MOBILE-->") != 1:
        sys.exit("deck_template.html harus memuat tepat satu penanda <!--MOBILE-->")
    (OUT_DIR / "deck.html").write_text(dt.replace("<!--MOBILE-->", mobile), encoding="utf-8")
    print(f"OK: kamus {len(entries)} entri | deck {len(groups)} kosakata, {ndays} hari | latihan {len(pool)} soal -> {OUT_DIR}")


if __name__ == "__main__":
    main()
